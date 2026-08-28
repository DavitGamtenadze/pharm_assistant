from pathlib import Path

import pymupdf

from pharm_assistant.services.extraction import (
    extract_document,
    extract_pdf,
    is_usable_text,
    sanitize_extracted_text,
    strip_running_chrome,
)

HEADER = "MEDICATION SAFETY IN HIGH-RISK SITUATIONS"


def _multi_page_pdf(tmp_path, texts: list[str]):
    document = pymupdf.open()
    for text in texts:
        page = document.new_page()
        page.insert_textbox(pymupdf.Rect(48, 48, 560, 760), text, fontsize=11)
    path = tmp_path / "guideline.pdf"
    path.write_bytes(document.tobytes())
    document.close()
    return path


def test_sanitize_extracted_text_removes_converter_artifacts() -> None:
    dirty = (
        "High Risk29Print > Introduction\n"
        "<!-- Start of picture text -->ignored<!-- End of picture text -->\n"
        "<u>LASA</u> names increase selection errors."
    )

    cleaned = sanitize_extracted_text(dirty)

    assert "29Print" not in cleaned
    assert "<!--" not in cleaned
    assert "<u>" not in cleaned
    assert "LASA names increase selection errors" in cleaned
    assert sanitize_extracted_text("High Risk29Print") == ""
    assert sanitize_extracted_text("**2. OVERVIEW**") == "2. OVERVIEW"


def test_strip_running_chrome_removes_repeated_headers_and_page_numbers() -> None:
    pages = [
        f"{HEADER}\nHigh-alert medicines require independent double checks.\n1 of 4",
        f"{HEADER}\nLook-alike sound-alike names increase selection errors.\n2 of 4",
        f"{HEADER}\nA Manual of Operations should define local safety checks.\n3 of 4",
        f"{HEADER}\nStaff training should include high-risk situations.\n4 of 4",
    ]

    cleaned = strip_running_chrome(pages)

    assert all("MEDICATION SAFETY" not in page for page in cleaned)
    assert all("of 4" not in page for page in cleaned)
    assert "independent double checks" in cleaned[0]
    assert "Look-alike sound-alike" in cleaned[1]


def test_extract_pdf_keeps_page_numbers_and_drops_repeated_headers(tmp_path) -> None:
    path = _multi_page_pdf(
        tmp_path,
        [
            f"{HEADER}\nCover page with no usable body text.\n1 of 4",
            f"{HEADER}\nHigh-alert medicines require independent double checks.\n2 of 4",
            f"{HEADER}\nLook-alike sound-alike names increase selection errors.\n3 of 4",
            f"{HEADER}\nLocal procedures should define who performs the check.\n4 of 4",
        ],
    )

    pages = extract_pdf(path)

    assert [page.page_number for page in pages] == [1, 2, 3, 4]
    assert all(HEADER not in page.text for page in pages)
    assert "independent double checks" in pages[1].text
    assert is_usable_text(pages[1].text)


def test_extract_docx_and_plain_text(tmp_path: Path) -> None:
    from docx import Document

    path = tmp_path / "monograph.docx"
    document = Document()
    document.add_heading("Dosage and administration", level=1)
    document.add_paragraph(
        "High-alert medicines require an independent double check before administration."
    )
    document.save(path)

    pages = extract_document(path)
    assert pages[0].page_number == 1
    assert "independent double check" in pages[0].text

    notes = tmp_path / "notes.txt"
    notes.write_text(
        "Look-alike sound-alike names increase selection errors in busy clinics.",
        encoding="utf-8",
    )
    text_pages = extract_document(notes)
    assert "Look-alike sound-alike" in text_pages[0].text

    rtf = tmp_path / "note.rtf"
    rtf.write_text(
        r"{\rtf1\ansi High-alert medicines require an independent double check.}",
        encoding="utf-8",
    )
    assert "independent double check" in extract_document(rtf)[0].text


def test_extract_scanned_pdf_uses_ocr(tmp_path: Path) -> None:
    source = pymupdf.open()
    page = source.new_page()
    page.insert_text(
        (72, 140),
        "High-alert medicines require an independent double check.",
        fontsize=16,
    )
    pixmap = page.get_pixmap(matrix=pymupdf.Matrix(2, 2), alpha=False)
    source.close()

    scanned = pymupdf.open()
    image_page = scanned.new_page()
    image_page.insert_image(image_page.rect, pixmap=pixmap)
    path = tmp_path / "scan.pdf"
    path.write_bytes(scanned.tobytes())
    scanned.close()

    pages = extract_pdf(path)
    assert any("independent double check" in page.text.lower() for page in pages)
